"""Word 文档处理器 - 按段落提取 + 递归字符分片

分片流程:
1. 使用 python-docx 逐段落提取文本
2. 用换行符连接段落，保留段落边界
3. 使用 RecursiveCharacterTextSplitter 按字符递归分割

依赖: python-docx (>=1.1.0) - MIT 协议，Word 文档读取的事实标准库

不支持:
- .doc 格式（旧版 Word 97-2003 二进制格式）
- 如需支持 .doc，可考虑先用 LibreOffice 转换为 .docx 或使用 textract

选型理由:
- python-docx: 轻量、纯 Python、API 简洁
- textract: 功能全但依赖过多（需要安装 LibreOffice 等）
"""

from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from loguru import logger

from app.services.processors.base import DocumentProcessor


class WordProcessor(DocumentProcessor):
    """Word 文档处理器（仅支持 .docx）"""

    supported_extensions = {"docx"}

    def __init__(self):
        super().__init__()
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size * 2,
            chunk_overlap=self.chunk_overlap,
            length_function=len,
            is_separator_regex=False,
        )
        logger.info(
            f"WordProcessor 初始化完成, chunk_size={self.chunk_size * 2}, "
            f"overlap={self.chunk_overlap}"
        )

    def extract_text(self, file_path: str) -> str:
        """使用 python-docx 提取 Word 文档文本

        逐段落读取，保留段落间的换行分隔。
        支持:
        - 普通段落文本
        - 表格中的文本（逐行提取，以 | 分隔单元格）
        """
        from docx import Document as DocxDocument

        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        doc = DocxDocument(str(path))
        paragraphs = []

        # 1. 提取所有段落文本
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # 2. 提取表格中的文本
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(c for c in cells if c)
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        logger.info(
            f"Word 文本提取完成: {file_path}, "
            f"段落数: {len(doc.paragraphs)}, 表格数: {len(doc.tables)}, "
            f"文本长度: {len(full_text)} 字符"
        )
        return full_text

    def split(self, content: str, file_path: str = "") -> List[Document]:
        """
        Word 文档分片策略:

        使用 RecursiveCharacterTextSplitter 按字符递归分割。
        分割器优先级: "\\n\\n"（段落边界）> "\\n" > " " > ""
        这意味着分割优先发生在段落之间，保持语义完整性。
        """
        if not content or not content.strip():
            logger.warning(f"Word 文档内容为空: {file_path}")
            return []

        docs = self.text_splitter.create_documents(
            texts=[content],
            metadatas=[self._build_metadata(file_path)],
        )
        logger.info(f"Word 分割完成: {file_path} -> {len(docs)} 个分片")
        return docs
